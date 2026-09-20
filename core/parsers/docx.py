"""Word SI and BL parser.  OWNER: Sheng Kuan.

8 attachments, all table-based, all with BILINGUAL labels of the form
"Port of Loading (<chinese>)" and "Gross Weight (<chinese> KGS)".
aliases.normalise_label already strips the CJK - use it, do not re-implement.

THE TRAP: one bilingual label breaks the alias lookup outright. Verified on
email_055_BL.docx: 'Gross Wt (kgs) (毛重 KGS)' fails field_for_label because
the English alias itself ends in a bracket group ("gross wt (kgs)"), so once
normalise_label strips the CJK you are left with a DOUBLED bracket group -
"gross wt (kgs) (kgs)" - which matches nothing in the alias table. Fixed
locally in _resolve_field() by stripping the trailing "(...)" group and
retrying, at most twice. Nothing in aliases.py changes.
"""
from __future__ import annotations

import re

import docx

from core import aliases
from core.parsers import detect_kind, register
from core.types import CompareField, ExtractedDoc, FieldValue

#: Matches one trailing, balanced (non-nested) "(...)" group at the end of a
#: label, e.g. the " (毛重 KGS)" in "Gross Wt (kgs) (毛重 KGS)".
_TRAILING_BRACKET = re.compile(r"\s*\([^()]*\)\s*$")


def _resolve_field(label: str) -> "CompareField | None":
    """field_for_label with retries that strip a trailing bracket group.

    Confirmed: field_for_label('Gross Wt (kgs) (毛重 KGS)') returns None on
    the first try. Stripping the trailing "(...)" group once yields
    'Gross Wt (kgs)', which resolves to gross_weight_kg. Capped at two
    retries so a pathological label cannot loop.
    """
    candidate = label
    field = aliases.field_for_label(candidate)
    tries = 0
    while field is None and tries < 2:
        stripped = _TRAILING_BRACKET.sub("", candidate).strip()
        if not stripped or stripped == candidate:
            break
        candidate = stripped
        field = aliases.field_for_label(candidate)
        tries += 1
    return field


class DocxParser:
    extensions = (".docx",)

    def parse(self, path: str) -> ExtractedDoc:
        try:
            document = docx.Document(path)
        except Exception as exc:
            return ExtractedDoc(path=path, kind="UNREADABLE", readable=False, error=str(exc))

        try:
            title = ""
            for para in document.paragraphs:
                if para.text and para.text.strip():
                    title = para.text.strip()
                    break
            kind = detect_kind(title)

            fields: dict[CompareField, FieldValue] = {}
            text_lines: list[str] = []

            for table in document.tables:
                for row in table.rows:
                    cells = row.cells
                    if len(cells) < 2:
                        continue
                    label = (cells[0].text or "").strip()
                    value = (cells[1].text or "").strip()
                    if not label:
                        continue

                    text_lines.append(f"{label}: {value}")

                    field = _resolve_field(label)
                    if field is None or field in fields:
                        # Unknown label, or the field is already captured -
                        # keep the FIRST occurrence, not the last.
                        continue

                    # A blank token ("???", "TBA", ...) means the row exists
                    # but was left empty. Store "", never omit the field: a
                    # missing key means "no such row", a different
                    # escalation reason from "present but blank".
                    stored_value = "" if aliases.is_blank(value) else value

                    fields[field] = FieldValue(
                        value=stored_value,
                        raw=f"{label}: {value}",
                        line_no=0,  # position is not meaningful for docx cells
                        label=label,
                        decided_by="rule",
                    )

            text = "\n".join(text_lines)
            return ExtractedDoc(path=path, kind=kind, fields=fields, text=text, readable=True)
        except Exception as exc:
            return ExtractedDoc(path=path, kind="UNREADABLE", readable=False, error=str(exc))


register(DocxParser())
